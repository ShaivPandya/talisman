"""Census MARTS vintage-gate tests (LON-5 / DR-03 / DR-04)."""

from __future__ import annotations

import csv
import hashlib
import io
from dataclasses import asdict

import pytest
from openpyxl import load_workbook

from longaeva_app.collect.census_sources import (
    CENSUS_DIR,
    MANIFEST_PATH,
    SOURCES_DIR,
    latest_release_at_or_before,
    load_calendar,
    load_manifest,
)
from longaeva_app.collect.edgar_index import ORIGINS_CSV_PATH, PACKAGE_ROOT, SNAPSHOT_PATH, build_origins, load_snapshot
from longaeva_app.extract.census_marts import (
    CSV_COLUMNS,
    ObservationRow,
    ReleaseMeta,
    normalize_label,
    parse_cell,
    parse_marts_pdf,
    parse_release_line,
    parse_release_meta,
    sa_advance_total,
    series_key_for,
    split_label_and_values,
    validate_table1_header,
)

ADV2406_PDF = SOURCES_DIR / "adv2406.pdf"
ADV2506_PDF = SOURCES_DIR / "adv2506.pdf"
ADV2406_CSV = CENSUS_DIR / "adv2406.csv"
ADV2506_CSV = CENSUS_DIR / "adv2506.csv"

ParsedRelease = tuple[ReleaseMeta, list[ObservationRow]]


@pytest.fixture(scope="module")
def manifest() -> dict[str, object]:
    return load_manifest(MANIFEST_PATH)


@pytest.fixture(scope="module")
def calendar() -> list[dict[str, str]]:
    return load_calendar()


@pytest.fixture(scope="module")
def parsed_2406() -> ParsedRelease:
    return parse_marts_pdf(ADV2406_PDF, release_id="adv2406")


@pytest.fixture(scope="module")
def parsed_2506() -> ParsedRelease:
    return parse_marts_pdf(ADV2506_PDF, release_id="adv2506")


def test_manifest_hashes_and_timestamp_order(manifest: dict[str, object]) -> None:
    sources = manifest["sources"]
    assert isinstance(sources, list)
    assert len(sources) == 3
    for entry in sources:
        assert isinstance(entry, dict)
        path = PACKAGE_ROOT / str(entry["path"])
        assert path.is_file(), path
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        assert digest == entry["content_sha256"]
        if entry["kind"] == "advance_pdf":
            assert entry["publication_ts"] < entry["retrieval_ts"]


def test_release_times_edt_and_est() -> None:
    meta_a = parse_release_meta(ADV2406_PDF, release_id="adv2406")
    meta_b = parse_release_meta(ADV2506_PDF, release_id="adv2506")
    assert meta_a.publication_ts == "2024-07-16T12:30:00Z"
    assert meta_b.publication_ts == "2025-07-17T12:30:00Z"
    # EST case from the 2025 shutdown release (November).
    est_line = "FOR RELEASE AT 8:30 AM EST, TUESDAY, NOVEMBER 25, 2025\n"
    assert parse_release_line(est_line).strftime("%Y-%m-%dT%H:%M:%SZ") == "2025-11-25T13:30:00Z"
    # Legacy 2016 layout.
    legacy = (
        "FOR IMMEDIATE RELEASE\n"
        "WEDNESDAY, DECEMBER 14, 2016, AT 8:30 A.M. EST\n"
        "Rebecca DeNale CB16-208\n"
        "ADVANCE MONTHLY SALES FOR RETAIL AND FOOD SERVICES\n"
        "NOVEMBER 2016\n"
    )
    assert parse_release_line(legacy).strftime("%Y-%m-%dT%H:%M:%SZ") == "2016-12-14T13:30:00Z"


def test_committed_csvs_match_parser(parsed_2406: ParsedRelease, parsed_2506: ParsedRelease) -> None:
    for path, parsed in ((ADV2406_CSV, parsed_2406), (ADV2506_CSV, parsed_2506)):
        _meta, rows = parsed
        assert all(isinstance(r, ObservationRow) for r in rows)
        committed = path.read_text(encoding="utf-8")
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=CSV_COLUMNS, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow(asdict(row))
        assert buf.getvalue() == committed


def test_sa_totals_match_headlines(parsed_2406: ParsedRelease, parsed_2506: ParsedRelease) -> None:
    meta_a, rows_a = parsed_2406
    meta_b, rows_b = parsed_2506
    assert meta_a.headline_billions == 704.3
    assert meta_b.headline_billions == 720.1
    assert sa_advance_total(rows_a) == 704324.0
    assert sa_advance_total(rows_b) == 720106.0
    # Headline billions round from $ millions.
    assert round(sa_advance_total(rows_a) / 1000, 1) == meta_a.headline_billions
    assert round(sa_advance_total(rows_b) / 1000, 1) == meta_b.headline_billions


def _level(rows: list[ObservationRow], series: str, status: str, basis: str = "sa") -> float:
    for row in rows:
        if (
            row.table == "1"
            and row.series_key == series
            and row.measure == "level"
            and row.basis == basis
            and row.estimate_status == status
            and row.value
        ):
            return float(row.value)
    raise KeyError(series, status, basis)


def _pct(rows: list[ObservationRow], series: str, measure: str, status: str) -> float:
    for row in rows:
        if (
            row.table == "2"
            and row.series_key == series
            and row.measure == measure
            and row.estimate_status == status
            and row.value
        ):
            return float(row.value)
    raise KeyError(series, measure, status)


@pytest.mark.parametrize("series", ["retail_food_services_total", "retail_total", "naics_454", "naics_722"])
def test_table2_matches_table1_within_rounding(parsed_2406: ParsedRelease, series: str) -> None:
    _meta, rows = parsed_2406
    advance = _level(rows, series, "advance")
    prelim = _level(rows, series, "preliminary")
    mom = 100.0 * (advance - prelim) / prelim
    assert abs(mom - _pct(rows, series, "mom_pct", "advance")) < 0.05


def test_revision_example_june_2024(parsed_2406: ParsedRelease, parsed_2506: ParsedRelease) -> None:
    _a, rows_a = parsed_2406
    _b, rows_b = parsed_2506
    first = _level(rows_a, "retail_food_services_total", "advance")
    # In adv2506 the June 2024 SA total appears as the year-ago column.
    revised = None
    for row in rows_b:
        if (
            row.table == "1"
            and row.series_key == "retail_food_services_total"
            and row.basis == "sa"
            and row.measure == "level"
            and row.period_start == "2024-06-01"
            and row.value
        ):
            revised = float(row.value)
            break
    assert first == 704324.0
    assert revised == 692922.0
    assert revised != first


def test_xlsx_is_not_a_vintage_source(manifest: dict[str, object], parsed_2406: ParsedRelease) -> None:
    sources = manifest["sources"]
    assert isinstance(sources, list)
    xlsx_entry = next(s for s in sources if isinstance(s, dict) and s.get("kind") == "revised_xlsx")
    path = PACKAGE_ROOT / str(xlsx_entry["path"])
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["2024"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[4]
    june_idx = next(i for i, c in enumerate(header) if c and "Jun" in str(c))
    adjusted = False
    xlsx_sa: float | None = None
    for row in rows:
        label = str(row[1] or "")
        if label.upper().startswith("ADJUSTED"):
            adjusted = True
            continue
        if label.upper().startswith("NOT ADJUSTED"):
            adjusted = False
            continue
        if adjusted and label == "Retail and food services sales, total":
            cell = row[june_idx]
            assert isinstance(cell, (int, float))
            xlsx_sa = float(cell)
            break
    assert xlsx_sa == 666040.0
    first = sa_advance_total(parsed_2406[1])
    assert xlsx_sa != first
    # Header cites the annual survey — not an advance vintage.
    assert any(row[0] and "Annual" in str(row[0]) for row in rows[:3])


def test_shared_series_and_department_code_change(parsed_2406: ParsedRelease, parsed_2506: ParsedRelease) -> None:
    keys_a = {r.series_key for r in parsed_2406[1] if r.table == "1"}
    keys_b = {r.series_key for r in parsed_2506[1] if r.table == "1"}
    for key in ("retail_food_services_total", "retail_total", "naics_454", "naics_722", "naics_452_dept"):
        assert key in keys_a
        assert key in keys_b
    code_a = {r.naics_code for r in parsed_2406[1] if r.series_key == "naics_452_dept" and r.table == "1"}
    code_b = {r.naics_code for r in parsed_2506[1] if r.series_key == "naics_452_dept" and r.table == "1"}
    assert code_a == {"4521"}
    assert code_b == {"4522"}


def test_calendar_covers_window(calendar: list[dict[str, str]]) -> None:
    assert len(calendar) == 116
    assert calendar[0]["release_id"] == "adv1611"
    assert calendar[-1]["release_id"] == "adv2606"
    assert all(r["publication_ts"] for r in calendar)
    assert all(r["integrity_flag"] in {"ok", "possibly_replaced"} for r in calendar)
    adv1812 = next(r for r in calendar if r["release_id"] == "adv1812")
    assert adv1812["integrity_flag"] == "possibly_replaced"
    assert float(adv1812["last_modified_gap_days"]) > 2.0


def test_origins_census_timing(calendar: list[dict[str, str]]) -> None:
    assert SNAPSHOT_PATH.exists()
    rows = build_origins(load_snapshot())
    assert all(r["census_status"] == "eligible" for r in rows)
    by_date = {r["cutoff_utc"][:10]: r for r in rows}

    oct_2025 = by_date["2025-10-28"]
    assert oct_2025["census_release"] == "adv2508"
    assert oct_2025["census_reference_month"] == "2025-08"
    assert oct_2025["census_publication_utc"] <= "2025-10-28T20:06:03Z"

    jan_2019 = by_date["2019-01-30"]
    assert jan_2019["census_release"] == "adv1811"
    assert jan_2019["census_reference_month"] == "2018-11"

    # No origin uses a release published after its cutoff.
    for row in rows:
        hit = latest_release_at_or_before(row["cutoff_utc"], calendar)
        assert hit is not None
        assert hit["release_id"] == row["census_release"]
        assert hit["publication_ts"] <= row["cutoff_utc"]

    # Committed CSV matches rebuild.
    committed = list(csv.DictReader(ORIGINS_CSV_PATH.open(encoding="utf-8")))
    assert [r["census_release"] for r in committed] == [r["census_release"] for r in rows]


def test_parser_unit_cases() -> None:
    label, values = split_label_and_values(
        "GAFO4 (*) (*) (*) 134,697 121,802 127,859 130,729 (*) 130,975 130,367 128,802 128,712"
    )
    assert label == "GAFO"
    assert len(values) == 12
    assert values[0] == "(*)"

    wrapped = split_label_and_values("total ……………………………….………..………4,138,967 2.8 703,632")
    assert "total" in wrapped[0].lower()
    assert wrapped[1][0] == "4,138,967"

    assert parse_cell("‐7.0") == ("-7.0", "")
    assert parse_cell("(*)") == ("", "unavailable")
    assert parse_cell("(S)") == ("", "suppressed")
    assert series_key_for("", "Total (excl. motor vehicle & parts)") == "total_ex_motor"
    assert normalize_label("Retail & food services, total") == "retail and food services, total"

    with pytest.raises(ValueError, match="disagree"):
        # Wrong months for a June reference (should be Jun May Apr Jun May ×2).
        validate_table1_header("Jan. Feb. Mar. Jan. Feb. Jan. Feb. Mar. Jan. Feb.", "2024-06")
